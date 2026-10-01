"""Build the Excel management report: one printable page that follows a Year cell.

Stage 5 of the pipeline (the Excel report, beside the Power BI dashboard).
Input:  data/powerbi/*.parquet (run pipeline/export.py first)
Output: reports/excel/saudi_spending_report.xlsx

The workbook is a working Excel file, not a picture of numbers. The Data sheet holds every row; the Calc sheet
works out each figure with ordinary Excel formulas (SUMIFS, MAXIFS, INDEX/MATCH); the Report sheet only shows
them. A manager picks a year in one cell and the whole page, the tables and the charts follow it.
The figures match the dashboard's: growth compares only the week numbers found in both years, and an average
payment is total spent divided by total payments.

The page looks like the dashboard (redesign of 1 Oct 2026): the same slate band with the title and the Year box,
white cards on a light grey page, the four headline cards with their weekly line, green ▲ and red ▼ for changes
against the same weeks a year before, and the same colours for the week types.
"""
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, NumDataSource, NumRef, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_FOLDER = Path(__file__).resolve().parent.parent
EXPORT_FOLDER = PROJECT_FOLDER / "data" / "powerbi"
OUTPUT = PROJECT_FOLDER / "reports" / "excel" / "saudi_spending_report.xlsx"

YEARS = [2022, 2023, 2024, 2025]  # 2021 is left out: its year before, 2020, starts in May, mid-pandemic
DEFAULT_YEAR = 2025

# The dashboard's colours, so the two reports look like one family
INK, INK2, MUTED = "13161B", "4F5966", "646E7B"
GRID, BORDER, CANVAS = "EDF0F3", "E3E7EC", "F4F6F8"
FRAME, QUIET, LIGHT_TEXT = "1D2939", "98A2B3", "D0D5DD"  # the slate band and its two text greys
UP, DOWN = "15803D", "C81E1E"  # changes against the same weeks a year before, always with ▲ / ▼
FOCUS, AXIS_LINE = "344054", "AEB7C2"
WEEK_COLOURS = {"Normal": "D0D5DD", "Part Ramadan": "5FB48C", "Full Ramadan": "16825D", "Eid": "D98A1E"}
YEAR_BEFORE = "5B6573"  # the dashed line: the same week a year before
CARD_LABEL, CARD_TEXT = "B5C9C0", "F3F6F4"  # text on the slate 'What to plan' box, as on the dashboard's Ramadan card
FONT = "Segoe UI"
# Growth on the Calc sheet: a sign always shown, falls in red. The true minus sign (−) is escaped with a backslash,
# because Excel only shows a few characters (+ - $ and so on) as they are inside a number format.
PLUS = '+0.0%;[Red]\\−0.0%;0.0%'
# A change on the Report: '▲ +6.4%' or '▼ −5.0%' (the colour comes from conditional formatting). A custom format
# for negatives shows the number without its own minus, so the section writes '▼ −' itself. The cells hold the change
# rounded to 0.1%, so a change that rounds to nothing reads '0.0%' with no arrow.
CHANGE = '"▲ +"0.0%;"▼ −"0.0%;0.0%'

# The page grid in screen pixels at 100% zoom (1,280 px wide, like the dashboard): four blocks of 300 px with
# 16 px gutters. Columns A to U, then the row heights of rows 1 to 53. In the second and fourth blocks the last two
# columns are 78 px: the header 'Avg. payment' is 70 px wide in bold 8 pt and wrapped onto a third line in 74 px.
COLUMN_PX = [16, 82, 62, 78, 78, 16, 76, 68, 78, 78, 16, 82, 62, 78, 78, 16, 76, 68, 78, 78, 16]
ROW_PX = ([10, 30, 18, 10, 16, 10, 18, 34, 20, 18, 10, 16, 10, 22, 18] + [20] * 12 + [10, 16, 10, 22, 18, 30]
          + [20] * 17 + [22, 10, 16])

# One row per week per series, the three levels stacked, with the week labels the dashboard uses.
DATA_QUERY = """
    WITH rows AS (
        SELECT 'National' AS level, 'Saudi Arabia' AS name, date_key, value_k_sar, transactions_k
        FROM '{folder}/fact_national.parquet'
        UNION ALL
        SELECT 'City', c.city, f.date_key, f.value_k_sar, f.transactions_k
        FROM '{folder}/fact_city.parquet' AS f JOIN '{folder}/dim_city.parquet' AS c USING (city_key)
        UNION ALL
        SELECT 'Sector', s.sector, f.date_key, f.value_k_sar, f.transactions_k
        FROM '{folder}/fact_sector.parquet' AS f JOIN '{folder}/dim_sector.parquet' AS s USING (sector_key)
    )
    SELECT d.week_start, d.year, d.week_of_year, r.level, r.name, r.value_k_sar, r.transactions_k,
           d.week_type, coalesce(d.phase, '') AS phase, coalesce(d.ramadan_phase, '') AS ramadan_phase,
           d.ramadan_year,
           -- one text key per row, so a formula can find a week with a single MATCH
           r.level || '|' || r.name || '|' || d.year || '|' || d.week_of_year AS key
    FROM rows AS r JOIN '{folder}/dim_date.parquet' AS d USING (date_key)
    ORDER BY CASE r.level WHEN 'National' THEN 1 WHEN 'City' THEN 2 ELSE 3 END, r.name, d.week_start
"""
DATA_COLUMNS = [  # (header on the Data sheet, defined name the formulas use, column width, number format)
    ("Week start", "WeekStart", 12, "d mmm yyyy"),
    ("Year", "Year", 7, "0"),
    ("Week", "Week", 7, "0"),
    ("Level", "Level", 10, "@"),
    ("Name", "Name", 22, "@"),
    ("Spent (thousand SAR)", "Spent", 20, "#,##0"),
    ("Payments (thousand)", "Payments", 19, "#,##0"),
    ("Week type", "WeekType", 13, "@"),
    ("Phase", "Phase", 30, "@"),
    ("Ramadan phase", "RamadanPhase", 16, "@"),
    ("Ramadan year", "RamadanYear", 14, "0"),
    ("Key", "Key", 38, "@"),
]


def font(size=10, bold=False, colour=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=colour, italic=italic)


def fill(colour):
    return PatternFill("solid", start_color=colour, end_color=colour)


def side(colour):
    return Side(style="thin", color=colour)


def name_range(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def read_data():
    folder = EXPORT_FOLDER.as_posix()
    if not (EXPORT_FOLDER / "fact_national.parquet").exists():
        raise FileNotFoundError(f"{EXPORT_FOLDER} has no export. Run python pipeline/export.py first.")
    rows = duckdb.sql(DATA_QUERY.format(folder=folder)).fetchall()
    cities = [r[0] for r in duckdb.sql(f"SELECT city FROM '{folder}/dim_city.parquet' ORDER BY city").fetchall()]
    sectors = [r[0] for r in duckdb.sql(f"SELECT sector FROM '{folder}/dim_sector.parquet' ORDER BY sector").fetchall()]
    return rows, cities, sectors


def data_sheet(wb, rows):
    ws = wb.create_sheet("Data")
    ws.append([c[0] for c in DATA_COLUMNS])
    for row in rows:
        ws.append(list(row))
    last = len(rows) + 1
    for i, (header, name, width, number_format) in enumerate(DATA_COLUMNS):
        letter = ws.cell(row=1, column=i + 1).column_letter
        ws.column_dimensions[letter].width = width
        for cell in ws[letter][1:]:
            cell.number_format = number_format
        # Readable formulas: SUMIFS(Spent, Level, "City", ...) instead of SUMIFS(Data!$F$2:$F$7831, ...)
        name_range(wb, name, f"Data!${letter}$2:${letter}${last}")
    table = Table(displayName="Spending", ref=f"A1:{ws.cell(row=1, column=len(DATA_COLUMNS)).column_letter}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    return last - 1


def calc_sheet(wb, cities, sectors):
    """Every figure of the report, worked out in one place. Report cells only point here."""
    ws = wb.create_sheet("Calc")
    ws.column_dimensions["A"].width = 44
    for k in range(2, 31):
        ws.column_dimensions[get_column_letter(k)].width = 14
    ws["A1"] = "Calculations for the Report sheet. Every cell here is a formula over the Data sheet."
    ws["A1"].font = font(11, bold=True)

    same_weeks = 'Week, ">="&CommonFrom, Week, "<="&CommonTo'
    scalars = [  # (defined name, label, formula, number format)
        ("PrevYear", "Year before", "=SelYear-1", "0"),
        ("FirstWeekY", "First week number of the year", '=_xlfn.MINIFS(Week, Level, "National", Year, SelYear)', "0"),
        ("LastWeekY", "Last week number of the year", '=_xlfn.MAXIFS(Week, Level, "National", Year, SelYear)', "0"),
        ("FirstWeekP", "First week number of the year before", '=_xlfn.MINIFS(Week, Level, "National", Year, PrevYear)', "0"),
        ("LastWeekP", "Last week number of the year before", '=_xlfn.MAXIFS(Week, Level, "National", Year, PrevYear)', "0"),
        # Growth is fair only on weeks found in both years: 2025 has weeks 1 to 27 so far, 2023 has a week 53.
        ("CommonFrom", "Same weeks: from week", "=MAX(FirstWeekY, FirstWeekP)", "0"),
        ("CommonTo", "Same weeks: to week", "=MIN(LastWeekY, LastWeekP)", "0"),
        ("WeeksCommon", "Same weeks: how many", "=CommonTo-CommonFrom+1", "0"),
        ("WeeksY", "Weeks in the year", '=COUNTIFS(Level, "National", Year, SelYear)', "0"),
        ("DateFrom", "First day", '=_xlfn.MINIFS(WeekStart, Level, "National", Year, SelYear)', "d mmm yyyy"),
        ("DateTo", "Last day", '=_xlfn.MAXIFS(WeekStart, Level, "National", Year, SelYear)+6', "d mmm yyyy"),
        ("NatSpent", "National spent, thousand SAR", '=SUMIFS(Spent, Level, "National", Year, SelYear)', "#,##0"),
        ("NatPay", "National payments, thousand", '=SUMIFS(Payments, Level, "National", Year, SelYear)', "#,##0"),
        ("NatSpentSame", "National spent on the same weeks", f'=SUMIFS(Spent, Level, "National", Year, SelYear, {same_weeks})', "#,##0"),
        ("NatSpentPrev", "National spent, same weeks a year before", f'=SUMIFS(Spent, Level, "National", Year, PrevYear, {same_weeks})', "#,##0"),
        ("NatPaySame", "National payments on the same weeks", f'=SUMIFS(Payments, Level, "National", Year, SelYear, {same_weeks})', "#,##0"),
        ("NatPayPrev", "National payments, same weeks a year before", f'=SUMIFS(Payments, Level, "National", Year, PrevYear, {same_weeks})', "#,##0"),
        ("NatGrowth", "Spending growth, same weeks", "=NatSpentSame/NatSpentPrev-1", PLUS),
        ("NatPayGrowth", "Growth in payments, same weeks", "=NatPaySame/NatPayPrev-1", PLUS),
        # An average payment is total spent / total payments, never an average of weekly averages.
        ("NatAvgPay", "Average payment, SAR", "=NatSpent/NatPay", "0.0"),
        ("NatAvgPaySame", "Average payment on the same weeks, SAR", "=NatSpentSame/NatPaySame", "0.0"),
        ("NatAvgPayPrev", "Average payment a year before, SAR", "=NatSpentPrev/NatPayPrev", "0.0"),
        ("NatAvgPayChange", "Change in the average payment", "=NatAvgPaySame/NatAvgPayPrev-1", PLUS),
        ("BiggestWeek", "Biggest week, thousand SAR", '=_xlfn.MAXIFS(Spent, Level, "National", Year, SelYear)', "#,##0"),
        ("BiggestRow", "Its row in the weekly table below", "=MATCH(BiggestWeek, WeeklySpent, 0)", "0"),
        ("BiggestDate", "Its first day", "=INDEX(WeeklyDate, BiggestRow)", "d mmm yyyy"),
        ("BiggestPhase", "What that week was", "=INDEX(WeeklyPhase, BiggestRow)", "@"),
        # The biggest of the year before's weeks that this year also has (2025's weeks 1 to 27 meet 2024's 1 to 27)
        ("BiggestPrev", "Biggest of the same weeks a year before, thousand SAR",
         '=_xlfn.MAXIFS(Spent, Level, "National", Year, PrevYear, Week, ">="&FirstWeekY, Week, "<="&LastWeekY)', "#,##0"),
        ("BiggestChange", "Biggest week against that week", "=BiggestWeek/BiggestPrev-1", PLUS),
        # Ramadan of the chosen year. The 8 weeks before Ramadan 2025 start in December 2024, so they carry
        # ramadan_year 2025, not the calendar year.
        ("RamBefore", "Ramadan: average week of the 8 weeks before, thousand SAR",
         '=AVERAGEIFS(Spent, Level, "National", RamadanYear, SelYear, RamadanPhase, "8 weeks before")', "#,##0"),
        ("RamLast", "Ramadan: last full week, thousand SAR",
         '=SUMIFS(Spent, Level, "National", RamadanYear, SelYear, RamadanPhase, "Last full week")', "#,##0"),
        ("RamEid", "Ramadan: Eid al-Fitr week, thousand SAR",
         '=SUMIFS(Spent, Level, "National", RamadanYear, SelYear, RamadanPhase, "Eid week")', "#,##0"),
        ("RamLift", "Last full week against the 8 weeks before", "=RamLast/RamBefore-1", PLUS),
        ("RamEidDrop", "Eid week against the last full week", "=RamEid/RamLast-1", PLUS),
        # The cards' little weekly lines: each line and its year-before line share one scale, from the lowest to
        # the highest week of the two (as in the dashboard's cards), so the values go from 0 to 1.
        ("SpkSpentLo", "Card line, spent: lowest week", "=MIN(MIN(WeeklySpent), MIN(WeeklySpentBefore))", "#,##0"),
        ("SpkSpentHi", "Card line, spent: highest week", "=MAX(MAX(WeeklySpent), MAX(WeeklySpentBefore))", "#,##0"),
        ("SpkPayLo", "Card line, payments: lowest week", "=MIN(MIN(WeeklyPay), MIN(WeeklyPayBefore))", "#,##0"),
        ("SpkPayHi", "Card line, payments: highest week", "=MAX(MAX(WeeklyPay), MAX(WeeklyPayBefore))", "#,##0"),
        ("SpkAvgLo", "Card line, average payment: lowest week", "=MIN(MIN(WeeklyAvg), MIN(WeeklyAvgBefore))", "0.0"),
        ("SpkAvgHi", "Card line, average payment: highest week", "=MAX(MAX(WeeklyAvg), MAX(WeeklyAvgBefore))", "0.0"),
    ]
    for i, (name, label, formula, number_format) in enumerate(scalars, start=3):
        ws.cell(row=i, column=1, value=label).font = font(10, colour=INK2)
        cell = ws.cell(row=i, column=2, value=formula)
        cell.number_format = number_format
        name_range(wb, name, f"Calc!$B${i}")

    # Weekly table: one row per week number, for the charts and the biggest-week note
    top = len(scalars) + 5
    heads = ["Week", "First day", "Key", "Spent, thousand SAR", "Week type", "Phase", "Label",           # A-G
             "Normal", "Part Ramadan", "Full Ramadan", "Eid", "Same week, year before",                 # H-L (chart, bn SAR)
             "Spent a year before", "Payments, thousand", "Payments a year before",                     # M-O
             "Average payment", "Average payment a year before", "Ramadan phase", "Value label",        # P-S
             "Line: spent", "Line: spent a year before", "Dot: last week",                              # T-V (cards' lines)
             "Line: payments", "Line: payments a year before", "Dot: last week",                        # W-Y
             "Line: average payment", "Line: average payment a year before", "Dot: last week",          # Z-AB
             "Dot: biggest week"]                                                                        # AC
    ws.cell(row=top - 1, column=1, value="Weekly, the chosen year (bn SAR in the chart columns)").font = font(11, bold=True)
    for j, head in enumerate(heads, start=1):
        c = ws.cell(row=top, column=j, value=head)
        c.font = font(10, bold=True)
    first, last = top + 1, top + 53
    for n in range(1, 54):
        r = top + n
        ws.cell(row=r, column=1, value=n)
        has = f'COUNTIFS(Level, "National", Year, SelYear, Week, A{r})'
        ws.cell(row=r, column=2, value=f'=IF({has}=0, "", _xlfn.MINIFS(WeekStart, Level, "National", Year, SelYear, Week, A{r}))').number_format = "d mmm yyyy"
        ws.cell(row=r, column=3, value=f'="National|Saudi Arabia|"&SelYear&"|"&A{r}')
        ws.cell(row=r, column=4, value=f'=IF(B{r}="", "", SUMIFS(Spent, Level, "National", Year, SelYear, Week, A{r}))').number_format = "#,##0"
        # &"" keeps an empty Data cell as empty text: a formula that reads an empty cell gets 0 otherwise
        ws.cell(row=r, column=5, value=f'=IF(B{r}="", "", INDEX(WeekType, MATCH(C{r}, Key, 0))&"")')
        ws.cell(row=r, column=6, value=f'=IF(B{r}="", "", INDEX(Phase, MATCH(C{r}, Key, 0))&"")')
        # axis label: the month's name on its first week (days 1 to 7), blank on the others
        ws.cell(row=r, column=7, value=f'=IF(B{r}="", "", IF(DAY(B{r})<=7, TEXT(B{r}, "mmm"), ""))')
        # One column per week type: a chart colours each column by type, so every week keeps its colour
        # whatever year is picked. #N/A (not 0) leaves no bar and no label.
        for k, week_type in enumerate(["Normal", "Part Ramadan", "Full Ramadan", "Eid"]):
            ws.cell(row=r, column=8 + k, value=f'=IF(E{r}="{week_type}", D{r}/1000000, NA())').number_format = "0.00"
        before = f'COUNTIFS(Level, "National", Year, PrevYear, Week, A{r})'
        ws.cell(row=r, column=12, value=f'=IF(M{r}="", NA(), M{r}/1000000)').number_format = "0.00"
        # The same week a year before, payments, and the average payments, as numbers or "" (MIN and MAX skip "")
        ws.cell(row=r, column=13, value=f'=IF(OR(B{r}="", {before}=0), "", SUMIFS(Spent, Level, "National", Year, PrevYear, Week, A{r}))').number_format = "#,##0"
        ws.cell(row=r, column=14, value=f'=IF(B{r}="", "", SUMIFS(Payments, Level, "National", Year, SelYear, Week, A{r}))').number_format = "#,##0"
        ws.cell(row=r, column=15, value=f'=IF(M{r}="", "", SUMIFS(Payments, Level, "National", Year, PrevYear, Week, A{r}))').number_format = "#,##0"
        ws.cell(row=r, column=16, value=f'=IF(B{r}="", "", D{r}/N{r})').number_format = "0.0"
        ws.cell(row=r, column=17, value=f'=IF(M{r}="", "", M{r}/O{r})').number_format = "0.0"
        ws.cell(row=r, column=18, value=f'=IF(B{r}="", "", INDEX(RamadanPhase, MATCH(C{r}, Key, 0))&"")')
        # The weekly chart writes the value of the biggest week over its column; an invisible series carries it.
        # (Not the Eid al-Fitr week's as well, as on the dashboard: that week always follows Ramadan's last full week,
        # a much taller column, which its label overhung in every year, and Excel cannot move one label aside. Its
        # value is in the Ramadan table beside the chart.)
        ws.cell(row=r, column=19, value=f'=IF(B{r}="", NA(), IF(D{r}=BiggestWeek, D{r}/1000000, NA()))').number_format = "0.00"
        # The cards' lines, from 0 (the lowest week of the two years) to 1 (the highest)
        for col, value, lo, hi in [(20, f"D{r}", "SpkSpentLo", "SpkSpentHi"), (21, f"M{r}", "SpkSpentLo", "SpkSpentHi"),
                                   (23, f"N{r}", "SpkPayLo", "SpkPayHi"), (24, f"O{r}", "SpkPayLo", "SpkPayHi"),
                                   (26, f"P{r}", "SpkAvgLo", "SpkAvgHi"), (27, f"Q{r}", "SpkAvgLo", "SpkAvgHi")]:
            ws.cell(row=r, column=col, value=f'=IF({value}="", NA(), ({value}-{lo})/({hi}-{lo}))').number_format = "0.000"
        # the dot: on the last week for three cards, on the biggest week for the Biggest week card
        for col, line in [(22, "T"), (25, "W"), (28, "Z")]:
            ws.cell(row=r, column=col, value=f'=IF(A{r}=WeeksY, {line}{r}, NA())').number_format = "0.000"
        ws.cell(row=r, column=29, value=f'=IF(A{r}=BiggestRow, T{r}, NA())').number_format = "0.000"
    for name, col in [("WeeklySpent", "D"), ("WeeklyDate", "B"), ("WeeklyPhase", "F"), ("WeeklySpentBefore", "M"),
                      ("WeeklyPay", "N"), ("WeeklyPayBefore", "O"), ("WeeklyAvg", "P"), ("WeeklyAvgBefore", "Q")]:
        name_range(wb, name, f"Calc!${col}${first}:${col}${last}")
    # Chart ranges that stop at the year's last week (weeks run 1, 2, 3 ... with no gaps from 2021 on)
    for name, col in [("ChartLabels", "G"), ("ChartNormal", "H"), ("ChartPart", "I"), ("ChartFull", "J"),
                      ("ChartEid", "K"), ("ChartBefore", "L"), ("ChartValueLabels", "S"),
                      ("LineSpent", "T"), ("LineSpentBefore", "U"), ("DotSpent", "V"),
                      ("LinePay", "W"), ("LinePayBefore", "X"), ("DotPay", "Y"),
                      ("LineAvg", "Z"), ("LineAvgBefore", "AA"), ("DotAvg", "AB"), ("DotBiggest", "AC")]:
        ws.defined_names[name] = DefinedName(
            name, attr_text=f"Calc!${col}${first}:INDEX(Calc!${col}${first}:${col}${last}, WeeksY)")

    # Cities and sectors: one row each, in name order; the Report sorts them with LARGE + INDEX/MATCH.
    # The cities' list also holds the whole country, so the Report can rank it among the cities.
    same = 'Week, ">="&CommonFrom, Week, "<="&CommonTo'
    columns = [  # (head, formula, number format); {who} picks the rows: Level, "City", Name, $A5 and so on
        ("Spent", 'SUMIFS(Spent, {who}, Year, SelYear)', "#,##0"),                                     # B
        ("Payments", 'SUMIFS(Payments, {who}, Year, SelYear)', "#,##0"),                               # C
        ("Spent, same weeks", f'SUMIFS(Spent, {{who}}, Year, SelYear, {same})', "#,##0"),              # D
        ("Spent, year before", f'SUMIFS(Spent, {{who}}, Year, PrevYear, {same})', "#,##0"),           # E
        ("Spending growth", "D{r}/E{r}-1", PLUS),                                                      # F
        ("Payments, same weeks", f'SUMIFS(Payments, {{who}}, Year, SelYear, {same})', "#,##0"),        # G
        ("Payments, year before", f'SUMIFS(Payments, {{who}}, Year, PrevYear, {same})', "#,##0"),     # H
        ("Payments growth", "G{r}/H{r}-1", PLUS),                                                      # I
        ("Share of national", "B{r}/NatSpent", "0.0%"),                                                 # J
        ("Average payment", "B{r}/C{r}", "0.0"),                                                        # K
        ("Average payment change", "(D{r}/G{r})/(E{r}/H{r})-1", PLUS),                                 # L
        # Sort key: the value plus a tiny row number, so two equal values never tie in MATCH
        ("Sort key", "{key}+ROW()/1000000000", "0.000000"),                                            # M
        ("Ramadan: 8 weeks before", 'AVERAGEIFS(Spent, {who}, RamadanYear, SelYear, RamadanPhase, "8 weeks before")', "#,##0"),  # N
        ("Ramadan: last full week", 'SUMIFS(Spent, {who}, RamadanYear, SelYear, RamadanPhase, "Last full week")', "#,##0"),      # O
        ("Ramadan lift", "O{r}/N{r}-1", PLUS),                                                         # P
        ("Lift sort key", "P{r}+ROW()/1000000000", "0.000000"),                                        # Q
    ]

    def block(top_row, title, first_head, items, key):
        """items: (name shown, how its rows are picked); key: the column the Report sorts by."""
        ws.cell(row=top_row - 1, column=1, value=title).font = font(11, bold=True)
        ws.cell(row=top_row, column=1, value=first_head).font = font(10, bold=True)
        for j, (head, _, _) in enumerate(columns, start=2):
            ws.cell(row=top_row, column=j, value=head).font = font(10, bold=True)
        for i, (shown, who) in enumerate(items, start=1):
            r = top_row + i
            ws.cell(row=r, column=1, value=shown)
            for j, (_, formula, number_format) in enumerate(columns, start=2):
                text = formula.format(r=r, who=who.format(r=r), key=f"{key}{r}")
                ws.cell(row=r, column=j, value="=" + text).number_format = number_format
        return top_row + 1, top_row + len(items)

    c_top = last + 4
    city_items = [(c, 'Level, "City", Name, $A{r}') for c in cities] + [("National", 'Level, "National"')]
    c_first, c_last = block(c_top, "Cities, and the whole country (ranked with them on the Report)", "City", city_items, "F")
    s_top = c_last + 4
    s_first, s_last = block(s_top, "Sectors", "Sector", [(s, 'Level, "Sector", Name, $A{r}') for s in sectors], "B")
    for prefix, a, b in [("City", c_first, c_last), ("Sector", s_first, s_last)]:
        for name, col in [("Name", "A"), ("Spent", "B"), ("Growth", "F"), ("PayGrowth", "I"), ("Share", "J"),
                          ("AvgPay", "K"), ("AvgPayChange", "L"), ("Key", "M"), ("Lift", "P"), ("LiftKey", "Q")]:
            name_range(wb, prefix + name, f"Calc!${col}${a}:${col}${b}")
    for row in ws.iter_rows(min_row=3, max_row=s_last):
        for cell in row:
            if cell.font == Font():  # not styled yet
                cell.font = font(10)
    return first, last


# ---------- the Report page ----------
def text_props(points, colour, bold=False, outline=None):
    """Font of a chart's text (axis labels, value labels); outline: a line of that width in points around each letter,
    in the text's own colour."""
    line = LineProperties(solidFill=colour, w=int(outline * 12700)) if outline else None
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(
        sz=int(points * 100), b=bold, solidFill=colour, ln=line, latin=DrawingFont(typeface=FONT))),
        endParaRPr=CharacterProperties())])


def no_frame(chart):
    """A chart with no background and no border of its own: the white card behind it shows through."""
    chart.graphical_properties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    chart.roundedCorners = False
    chart.legend = None


def line_style(series, colour, width_pt, dash=None):
    series.graphicalProperties.line.solidFill = colour
    series.graphicalProperties.line.width = int(width_pt * 12700)  # points to EMU
    if dash:
        series.graphicalProperties.line.dashStyle = dash
    series.smooth = False
    series.marker.symbol = "none"


def point_at(series, name):
    """The series reads a defined name that grows and shrinks with the year (27 weeks for 2025, 53 for 2023)."""
    series.val = NumDataSource(numRef=NumRef(f=f"Calc!{name}"))
    series.cat = AxDataSource(strRef=StrRef(f="Calc!ChartLabels"))


def place(ws, chart, start, end):
    """Pin a chart from one cell corner to another, each given as (column, row, x offset px, y offset px). Pinned to
    cells, a chart follows the columns as Excel draws them. (A chart given a size in pixels came out about 5% wider
    than its card in Excel for the web, whose columns are a little narrower than 7 px a character.)"""
    def marker(column, row, dx, dy):
        return AnchorMarker(col=column_index_from_string(column) - 1, colOff=pixels_to_EMU(dx), row=row - 1,
                            rowOff=pixels_to_EMU(dy))
    chart.anchor = TwoCellAnchor(_from=marker(*start), to=marker(*end))
    ws.add_chart(chart)


def report_sheet(wb, weekly_first, weekly_last, n_cities, n_sectors):
    ws = wb["Report"]
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False  # a page to read; the Calc and Data sheets keep their headings
    ws.sheet_view.zoomScale = 100
    calc_ws = wb["Calc"]
    # Column widths in Excel's units: a column of n characters of the default font (Calibri 11, 7 px a digit)
    # is 7n + 5 px wide. Row heights in points: 0.75 of a pixel.
    for k, px in enumerate(COLUMN_PX, start=1):
        ws.column_dimensions[get_column_letter(k)].width = round((px - 5) / 7, 3)
    for r, px in enumerate(ROW_PX, start=1):
        ws.row_dimensions[r].height = px * 0.75
    letters = [get_column_letter(k) for k in range(1, len(COLUMN_PX) + 1)]
    px_of = dict(zip(letters, COLUMN_PX))

    def cells(c1, r1, c2, r2):
        for r in range(r1, r2 + 1):
            for k in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                yield ws.cell(row=r, column=k)

    def paint(c1, r1, c2, r2, colour):
        for c in cells(c1, r1, c2, r2):
            c.fill = fill(colour)

    cards = []

    def card(c1, r1, c2, r2):
        """A white card; its thin grey frame is drawn at the end (see frame), after every merge."""
        paint(c1, r1, c2, r2, "FFFFFF")
        cards.append((c1, r1, c2, r2))

    def frame(c1, r1, c2, r2):
        """The card's frame, keeping the lines already inside it. Drawn last: merging cells copies the first cell's
        borders to the whole range, which wiped the frame on a merged edge (the legend's right end)."""
        k1, k2 = column_index_from_string(c1), column_index_from_string(c2)
        for c in cells(c1, r1, c2, r2):
            b = c.border
            c.border = Border(left=side(BORDER) if c.column == k1 else b.left,
                              right=side(BORDER) if c.column == k2 else b.right,
                              top=side(BORDER) if c.row == r1 else b.top,
                              bottom=side(BORDER) if c.row == r2 else b.bottom)

    def rule(c1, c2, r, colour, where="bottom"):
        """A line under (or over) a table row, across its columns, keeping a card's own frame."""
        for c in cells(c1, r, c2, r):
            b = c.border
            line = side(colour)
            c.border = Border(left=b.left, right=b.right, top=line if where == "top" else b.top,
                              bottom=line if where == "bottom" else b.bottom)

    def put(ref, value, size=9, bold=False, colour=INK, align="left", valign="center", indent=1, wrap=False,
            number_format=None, merge_to=None):
        c = ws[ref]
        c.value = value
        c.font = font(size, bold, colour)
        c.alignment = Alignment(horizontal=align, vertical=valign, indent=indent if align == "left" else 0,
                                wrap_text=wrap)
        if number_format:
            c.number_format = number_format
        if merge_to:
            ws.merge_cells(f"{ref}:{merge_to}")
        return c

    changes = []  # cells holding a change: green ▲ / red ▼ by conditional formatting

    def change(ref, formula, align="left", indent=1, size=9):
        changes.append(ref)
        return put(ref, f"=ROUND({formula}, 3)", size, True, INK2, align, indent=indent, number_format=CHANGE)

    # The light grey page, then the slate band across the top (the dashboard's frame)
    paint("A", 1, "U", len(ROW_PX), CANVAS)
    paint("A", 1, "U", 4, FRAME)
    put("B2", "Card spending in Saudi Arabia", 17, True, "FFFFFF", merge_to="L2")
    put("B3", "Management summary for planning teams in banks, shops and malls.   Weekly data to 12 Jul 2025.   "
              "Source: SAMA, via the KAPSARC data portal.", 9, colour=QUIET, merge_to="O3")
    # The one input cell: a white box in the band. Everything on the page follows it.
    put("Q2", "Year", 10, colour=LIGHT_TEXT, align="right", merge_to="R2")
    year = put("S2", DEFAULT_YEAR, 12, True, INK, "center", number_format="0", merge_to="T2")
    for c in cells("S", 2, "T", 2):
        c.fill = fill("FFFFFF")
    choice = DataValidation(type="list", formula1='"' + ",".join(map(str, YEARS)) + '"', allow_blank=False,
                            showErrorMessage=True, errorTitle="Year", error="Pick a year from 2022 to 2025.")
    ws.add_data_validation(choice)
    choice.add("S2")
    name_range(wb, "SelYear", "Report!$S$2")
    put("Q3", "Pick a year: the whole page follows it", 8, colour=QUIET, align="right", merge_to="T3")
    ws.sheet_view.selection[0].activeCell = "S2"  # the file opens on the Year box, its list arrow showing
    ws.sheet_view.selection[0].sqref = "S2"

    # Four headline cards: the number, its weekly line, the change against the same weeks a year before
    def versus(value, number_format, same, whole):
        """'vs 340.30 in 2024', or 'vs 557.31, 52 weeks of 2022' when the year has weeks the year before has not."""
        return (f'="vs "&TEXT({value}, "{number_format}")&IF(TEXT({same}, "{number_format}")<>TEXT({whole}, "{number_format}"), '
                f'", "&WeeksCommon&" weeks of "&PrevYear, " in "&PrevYear)')
    kpis = [
        ("B", "Spent (bn SAR)", "=NatSpent/1000000", "#,##0.00", "NatGrowth",
         versus("NatSpentPrev/1000000", "#,##0.00", "NatSpentSame/1000000", "NatSpent/1000000"),
         '=WeeksY&" weeks, "&IF(YEAR(DateFrom)=YEAR(DateTo), TEXT(DateFrom, "d mmm"), TEXT(DateFrom, "d mmm yyyy"))&" to "&TEXT(DateTo, "d mmm yyyy")',
         ("LineSpent", "LineSpentBefore", "DotSpent")),
        ("G", "Payments (millions)", "=NatPay/1000", "#,##0", "NatPayGrowth",
         versus("NatPayPrev/1000", "#,##0", "NatPaySame/1000", "NatPay/1000"), "Card payments at shop terminals",
         ("LinePay", "LinePayBefore", "DotPay")),
        ("L", "Average payment (SAR)", "=NatAvgPay", "#,##0.0", "NatAvgPayChange",
         versus("NatAvgPayPrev", "#,##0.0", "NatAvgPaySame", "NatAvgPay"), "Spent ÷ payments",
         ("LineAvg", "LineAvgBefore", "DotAvg")),
        ("Q", "Biggest week (bn SAR)", "=BiggestWeek/1000000", "#,##0.00", "BiggestChange",
         '="vs "&TEXT(BiggestPrev/1000000, "#,##0.00")&" in "&PrevYear',
         '=TEXT(BiggestDate, "d mmm")&IF(BiggestPhase="", "", ", "&BiggestPhase)',
         ("LineSpent", "LineSpentBefore", "DotBiggest")),
    ]
    for a, label, value, number_format, growth, compared, note, (line_name, before_name, dot_name) in kpis:
        k = letters.index(a)
        b2, b3, b4 = letters[k + 1], letters[k + 2], letters[k + 3]
        card(a, 6, b4, 11)
        put(f"{a}7", label, 9, colour=INK2)
        put(f"{a}8", value, 20, number_format=number_format, merge_to=f"{b2}8")
        change(f"{a}9", growth)
        put(f"{b2}9", compared, 9, colour=MUTED, indent=0)
        put(f"{a}10", note, 8, colour=MUTED)
        # The weekly line, as on the dashboard's cards: this year in grey, the same weeks a year before dotted
        # under it, and a dot on the last week (on the biggest week for the Biggest week card). The values run
        # from 0 to 1 (see SpkSpentLo on the Calc sheet), so the axis is fixed and hidden.
        spark = LineChart()
        for name in (before_name, line_name, dot_name):
            spark.add_data(Reference(calc_ws, min_col=20, min_row=weekly_first, max_row=weekly_last))
            point_at(spark.series[-1], name)
        line_style(spark.series[0], AXIS_LINE, 1.0, "sysDash")
        line_style(spark.series[1], INK2, 1.5)
        dot = spark.series[2]
        dot.graphicalProperties.line.noFill = True
        dot.marker = Marker(symbol="circle", size=5)
        dot.marker.graphicalProperties = GraphicalProperties(solidFill=FOCUS, ln=LineProperties(solidFill=FOCUS))
        spark.y_axis.scaling.min, spark.y_axis.scaling.max = -0.15, 1.15
        spark.x_axis.delete = spark.y_axis.delete = True
        spark.y_axis.majorGridlines = None
        spark.layout = Layout(manualLayout=ManualLayout(x=0, y=0, w=1, h=1, xMode="edge", yMode="edge"))
        no_frame(spark)
        place(ws, spark, (b3, 7, 6, 6), (b4, 8, px_of[b4] - 12, ROW_PX[7] - 4))

    # Weekly chart card (left) and the Ramadan card (right)
    card("B", 13, "J", 28)
    card("L", 13, "T", 28)
    put("B14", "Weekly spending", 11, True)
    put("B15", '=SelYear&", bn SAR"', 9, colour=INK2)
    # The legend is text in a cell, on the title's right like the dashboard's: Excel's own legend took space from
    # the chart and covered its axis.
    parts = []
    for week_type in ["Normal", "Part Ramadan", "Full Ramadan", "Eid"]:
        parts += [TextBlock(InlineFont(rFont=FONT, sz=8, color=WEEK_COLOURS[week_type]), "■ "),
                  TextBlock(InlineFont(rFont=FONT, sz=8, color=INK2), f"{week_type}    ")]
    parts += [TextBlock(InlineFont(rFont=FONT, sz=8, color=YEAR_BEFORE), "- - "),
              TextBlock(InlineFont(rFont=FONT, sz=8, color=INK2), "Same weeks, year before")]
    # D to J: the legend is 404 px wide in Segoe UI 8 pt (E to J held 394)
    put("D15", None, 8, colour=INK2, align="right", merge_to="J15")
    ws["D15"] = CellRichText(*parts)
    ws["D15"].alignment = Alignment(horizontal="right", vertical="center", indent=1)

    bars = BarChart()
    bars.type, bars.grouping, bars.overlap, bars.gapWidth = "col", "stacked", 100, 40
    for k, (week_type, name) in enumerate(zip(WEEK_COLOURS, ["ChartNormal", "ChartPart", "ChartFull", "ChartEid"])):
        bars.add_data(Reference(calc_ws, min_col=8 + k, min_row=weekly_first, max_row=weekly_last))
        s = bars.series[-1]
        s.graphicalProperties.solidFill = WEEK_COLOURS[week_type]
        s.graphicalProperties.line.noFill = True
        point_at(s, name)
    lines = LineChart()
    for name in ("ChartBefore", "ChartValueLabels", "ChartValueLabels"):
        lines.add_data(Reference(calc_ws, min_col=12, min_row=weekly_first, max_row=weekly_last))
        point_at(lines.series[-1], name)
    line_style(lines.series[0], YEAR_BEFORE, 1.5, "dash")
    # The value label: a series with no line and no marker, a value only on the biggest week (#N/A elsewhere), its
    # label written above the point, which is the top of that week's column. The label comes twice: first in white
    # with a thick white outline, then in ink on top, so the dashed line passes behind the digits (the dashboard's
    # 'halo'; seen working in Excel for the web).
    for labels, colour, outline in ((lines.series[1], "FFFFFF", 2.5), (lines.series[2], INK, None)):
        labels.graphicalProperties.line.noFill = True
        labels.marker.symbol = "none"
        labels.dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                                     showPercent=False, dLblPos="t", numFmt="0.00",
                                     txPr=text_props(8, colour, bold=True, outline=outline))
    bars += lines
    no_frame(bars)
    bars.y_axis.numFmt = "0"
    bars.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=GRID, w=9525)))
    bars.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    bars.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=AXIS_LINE, w=9525))
    bars.y_axis.majorTickMark = bars.x_axis.majorTickMark = "none"
    bars.y_axis.txPr = text_props(8, MUTED)
    bars.x_axis.txPr = text_props(8, MUTED)
    bars.x_axis.tickLblSkip = 1  # the labels are month names on each month's first week, blank on the others
    bars.x_axis.tickMarkSkip = 1
    bars.x_axis.delete = False  # openpyxl 3.1 hides axes unless told
    bars.y_axis.delete = False
    place(ws, bars, ("B", 16, 4, 0), ("K", 28, 0, 0))  # B16 to the end of J27: the card less its title rows

    put("L14", "Ramadan and Eid", 11, True)
    put("L15", '="Ramadan "&SelYear&", the whole country. The 8 weeks before it can start in the year before."', 9,
        colour=INK2)
    put("L16", "Week", 8, True)
    put("N16", "bn SAR", 8, True, align="center")
    put("O16", "Change", 8, True, align="center")
    rule("L", "O", 16, INK)
    # the dashboard's names for the three points; 'Last full week of Ramadan' (147 px) was cut by the number beside it
    rows = [("8 weeks before (avg.)", "=RamBefore/1000000", None),
            ("Last full week", "=RamLast/1000000", "RamLift"),
            ("Eid al-Fitr week", "=RamEid/1000000", "RamEidDrop")]
    for r, (label, value, lift) in enumerate(rows, start=17):
        put(f"L{r}", label)
        put(f"N{r}", value, align="center", number_format="#,##0.00")
        if lift:
            change(f"O{r}", lift, align="center", indent=0)
        rule("L", "O", r, GRID)
    put("L20", "Last full week: against the 8 weeks before.\nEid week: against the last full week.", 8, colour=MUTED,
        valign="top", wrap=True, merge_to="O21")
    put("Q16", "Sectors in the last full week", 8, True)
    put("T16", "Change", 8, True, align="center")
    rule("Q", "T", 16, INK)
    picks = ["LARGE(SectorLiftKey, 1)", "LARGE(SectorLiftKey, 2)", "LARGE(SectorLiftKey, 3)", "SMALL(SectorLiftKey, 1)"]
    for r, pick in enumerate(picks, start=17):
        put(f"Q{r}", f"=INDEX(SectorName, MATCH({pick}, SectorLiftKey, 0))")
        change(f"T{r}", f"INDEX(SectorLift, MATCH({pick}, SectorLiftKey, 0))", align="center", indent=0)
        rule("Q", "T", r, GRID)
    put("Q21", '=IF(T20<0, "The 3 biggest rises and the biggest fall.", "The 3 biggest rises and the smallest rise.")', 8,
        colour=MUTED)
    # What to plan: the dashboard's Ramadan card in a box, slate like the band, with the year's own numbers
    paint("L", 23, "T", 27, FRAME)
    put("L23", "WHAT TO PLAN", 9, True, CARD_LABEL, valign="bottom", merge_to="T23")
    put("L24", '="Have stock, staff and offers ready for the last full week of Ramadan, above all in "&LOWER(Q17)&" and "'
               '&LOWER(Q18)&IF(BiggestPhase="the last full week of Ramadan", ": it was the biggest week of "&SelYear&", "'
               '&TEXT(RamLift, "0%")&" above the weeks before Ramadan.", ": in "&SelYear&" it was "&TEXT(RamLift, "0%")'
               '&" above the weeks before Ramadan.")&" The Eid week itself is quieter: in every year from 2021 to 2025 it '
               'fell against the week before."', 9, colour=CARD_TEXT, valign="top", wrap=True, merge_to="T26")

    # Cities, ranked by growth (left) and sectors, largest first (right)
    card("B", 30, "J", 31 + 3 + n_cities + 1)  # the cities and the National row, then 20 px of room
    card("L", 30, "T", 31 + 3 + n_sectors + 1)  # the sectors and the total row, then 10 px
    put("B31", "Cities, ranked by growth", 11, True)
    put("B32", "On the same weeks a year before. Cities above the National row grew faster.", 9, colour=INK2)
    put("L31", "Sectors, largest first", 11, True)
    put("L32", "On the same weeks a year before. Spending = payments × average payment.", 9, colour=INK2)
    heads = ["Spent\n(bn SAR)", "Share of\nnational", None, "Spending\ngrowth", "Payments\ngrowth",
             "Avg. payment\n(SAR)", "Avg. payment\nchange"]

    def table(first, name_head, prefix, rows_shown, total=None):
        f = letters.index(first)
        cols = [letters[f + k] for k in (2, 3, 4, 5, 6, 7, 8)]
        last_col = letters[f + 8]
        put(f"{first}33", name_head, 8, True, valign="bottom")
        for col, head in zip(cols, heads):
            if head:
                put(f"{col}33", head, 8, True, align="center", valign="bottom", wrap=True)
        rule(first, last_col, 33, INK)
        pick = f"MATCH(LARGE({prefix}Key, {{k}}), {prefix}Key, 0)"
        for k in range(1, rows_shown + 1):
            r, p = 33 + k, pick.format(k=k)
            put(f"{first}{r}", f"=INDEX({prefix}Name, {p})")
            put(f"{cols[0]}{r}", f"=INDEX({prefix}Spent, {p})/1000000", align="center", number_format="#,##0.00")
            put(f"{cols[1]}{r}", f"=INDEX({prefix}Share, {p})", align="center", number_format="0.0%")
            change(f"{cols[3]}{r}", f"INDEX({prefix}Growth, {p})", align="center", indent=0)
            change(f"{cols[4]}{r}", f"INDEX({prefix}PayGrowth, {p})", align="center", indent=0)
            put(f"{cols[5]}{r}", f"=INDEX({prefix}AvgPay, {p})", align="center", number_format="#,##0.0")
            change(f"{cols[6]}{r}", f"INDEX({prefix}AvgPayChange, {p})", align="center", indent=0)
            rule(first, last_col, r, GRID)
        if total:
            r = 34 + rows_shown
            put(f"{first}{r}", total, bold=True)
            for col, value, number_format in [(cols[0], "=NatSpent/1000000", "#,##0.00"), (cols[1], "=1", "0.0%"),
                                              (cols[5], "=NatAvgPay", "#,##0.0")]:
                put(f"{col}{r}", value, bold=True, align="center", number_format=number_format)
            for col, value in [(cols[3], "NatGrowth"), (cols[4], "NatPayGrowth"), (cols[6], "NatAvgPayChange")]:
                change(f"{col}{r}", value, align="center", indent=0)
            rule(first, last_col, r, INK, where="top")
        return f"{first}34:{last_col}{33 + rows_shown}"

    city_rows = table("B", "City", "City", n_cities + 1)
    table("L", "Sector", "Sector", n_sectors, total="All sectors")
    # The whole country's row among the cities: grey and bold, wherever its growth ranks it
    # (bold only, no colour: with a colour this rule beat the green and red of the row's changes)
    ws.conditional_formatting.add(city_rows, FormulaRule(formula=['$B34="National"'], fill=fill(CANVAS),
                                                         font=Font(bold=True)))
    # Changes: green with ▲ when up, red with ▼ when down (the arrow is in the number format), grey at 0.0%
    where = " ".join(changes)
    ws.conditional_formatting.add(where, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, bold=True, color=UP)))
    ws.conditional_formatting.add(where, CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, bold=True, color=DOWN)))

    for rectangle in cards:
        frame(*rectangle)

    # Notes under the cities
    put("B48", "Notes", 9, True, INK2)
    put("B49", "Spending is the value of card payments at shop terminals: no cash, no online shopping, so growth mixes "
               "real spending, prices and the move to cards. Same weeks: each week against the week with the same "
               "number a year before. The source gives sectors for the whole country only.", 8, colour=MUTED,
        valign="top", wrap=True, merge_to="J52")

    # One landscape A4 page
    ws.print_area = f"A1:U{len(ROW_PX)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4


def about_sheet(wb, n_rows):
    ws = wb.create_sheet("About")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.showRowColHeaders = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 100
    ws.column_dimensions["D"].width = 2
    # the same slate band as the Report
    for r, height in [(1, 10), (2, 30), (3, 10)]:
        ws.row_dimensions[r].height = height * 0.75
        for col in "ABCD":
            ws[f"{col}{r}"].fill = fill(FRAME)
    ws["B2"] = "About this workbook"
    ws["B2"].font = font(16, bold=True, colour="FFFFFF")
    ws["B2"].alignment = Alignment(vertical="center")
    lines = [
        ("How to use it", "Pick a year in the white Year box at the top right of the Report sheet (2022 to 2025). "
                          "Every number, table and chart follows it. The Report sheet prints on one landscape A4 page."),
        ("Report", "The page for the manager: four headline cards with their weekly lines, the weekly chart, Ramadan "
                   "and Eid, cities ranked by growth and sectors by size."),
        ("Calc", "Every figure worked out with ordinary formulas (SUMIFS, MAXIFS, AVERAGEIFS, INDEX/MATCH, LARGE). "
                 "The Report sheet only points here, so each number can be traced to one formula."),
        ("Data", f"{n_rows:,} rows: one row per week for the country, each of the 11 cities and each of the 17 "
                 "sectors (10 May 2020 to 12 Jul 2025). Spent is in thousand SAR, payments in thousands."),
        ("", None),
        ("Spending", "The value of card payments at shop terminals, in Saudi riyals."),
        ("Payment", "One card payment at a shop terminal."),
        ("Average payment", "Spending divided by the number of payments."),
        ("Same weeks", "Each week against the week with the same number a year before. Weeks found in only one of "
                       "the two years are left out, so a part year (2025 so far) is compared with the same part."),
        ("Changes", "Green with ▲ when up, red with ▼ when down, always against the same weeks a year before."),
        ("Weekly lines", "In each headline card: the year's weeks in grey, the same weeks a year before dotted, on one "
                         "scale from the lowest to the highest week; the dot marks the last week, or the biggest week."),
        ("Last full week of Ramadan", "The last Sunday-to-Saturday week with all 7 days in Ramadan."),
        ("Week type", "Full Ramadan: all 7 days in Ramadan. Part Ramadan: some of them. Eid: a week that holds an "
                      "Eid day. Normal: no day of Ramadan and no Eid."),
        ("", None),
        ("Source", "SAMA weekly point-of-sale statistics, via the KAPSARC data portal. 270 weeks, 10 May 2020 to "
                   "12 Jul 2025. Ramadan and Eid dates: the Umm al-Qura calendar."),
        ("What it leaves out", "Cash and online shopping. Growth mixes real spending, prices and the move from cash "
                               "to cards."),
        ("Cities and sectors", "The source gives sectors for the whole country only, never inside one city. \"Other "
                               "cities\" holds every city it does not name, including Eastern Province cities such "
                               "as Al-Ahsa and Jubail."),
        ("Built by", "pipeline/excel_report.py in the project repository, from the same files as the Power BI "
                     "dashboard, so the two show the same numbers."),
    ]
    for i, (term, text) in enumerate(lines, start=5):
        a = ws.cell(row=i, column=2, value=term)
        a.font = font(10, bold=True)
        a.alignment = Alignment(vertical="top")
        if text:
            b = ws.cell(row=i, column=3, value=text)
            b.font = font(10, colour=INK2)
            b.alignment = Alignment(wrap_text=True, vertical="top")
            # Excel does not grow a row to fit wrapped text when it opens a file, so set the height:
            # about 105 characters per line in the 100-wide column, 14 points per line, and 6 points of air
            ws.row_dimensions[i].height = 14 * (len(text) // 105 + 1) + 6


def main():
    rows, cities, sectors = read_data()
    wb = Workbook()
    wb.active.title = "Report"
    n_rows = data_sheet(wb, rows)
    weekly_first, weekly_last = calc_sheet(wb, cities, sectors)
    report_sheet(wb, weekly_first, weekly_last, len(cities), len(sectors))
    about_sheet(wb, n_rows)
    wb.move_sheet("About", offset=-2)  # Report, About, Data, Calc
    # openpyxl writes formulas without their results; this tells Excel to calculate everything on opening.
    wb.calculation.fullCalcOnLoad = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"{OUTPUT.relative_to(PROJECT_FOLDER)}: {n_rows:,} data rows, {len(cities)} cities, {len(sectors)} sectors")


if __name__ == "__main__":
    main()
